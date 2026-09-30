from __future__ import annotations

import json
import threading
from pathlib import Path


class DocumentParser:
    CHUNK_SIZE = 500
    CHUNK_OVERLAP = 50

    def parse(self, file_path: str | Path) -> list[dict]:
        path = Path(file_path)
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            return self.parse_pdf(path)
        if suffix == ".docx":
            return self.parse_docx(path)
        if suffix in {".txt", ".md"}:
            return self.parse_txt(path)
        raise ValueError(f"不支持的文档格式: {suffix}；仅支持 PDF/TXT/MD/DOCX")

    def parse_pdf(self, file_path: str | Path) -> list[dict]:
        try:
            import pymupdf as fitz
        except ImportError as exc:
            raise RuntimeError("缺少 PyMuPDF") from exc
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(path)
        chunks: list[dict] = []
        with fitz.open(path) as document:
            for page_index in range(document.page_count):
                text = document.load_page(page_index).get_text("text")
                chunks.extend(
                    self.chunk_text(
                        text,
                        {"source": path.name, "path": str(path), "page": page_index + 1},
                    )
                )
        return chunks

    def parse_txt(self, file_path: str | Path) -> list[dict]:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(path)
        raw = path.read_bytes()
        text = None
        for encoding in ("utf-8-sig", "utf-8", "gb18030", "big5"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            text = raw.decode("utf-8", errors="replace")
        return self.chunk_text(
            text, {"source": path.name, "path": str(path), "page": None}
        )

    def parse_docx(self, file_path: str | Path) -> list[dict]:
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("缺少 python-docx") from exc
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(path)
        document = Document(path)
        parts = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]
        for table in document.tables:
            for row in table.rows:
                line = "\t".join(
                    cell.text.strip() for cell in row.cells if cell.text.strip()
                )
                if line:
                    parts.append(line)
        return self.chunk_text(
            "\n".join(parts),
            {"source": path.name, "path": str(path), "page": None},
        )

    def chunk_text(self, text: str, metadata: dict) -> list[dict]:
        cleaned = "\n".join(line.strip() for line in text.splitlines() if line.strip())
        if not cleaned:
            return []
        chunks = []
        start = 0
        index = 0
        while start < len(cleaned):
            end = min(len(cleaned), start + self.CHUNK_SIZE)
            if end < len(cleaned):
                boundary = max(
                    cleaned.rfind("\n", start + self.CHUNK_SIZE // 2, end),
                    cleaned.rfind("。", start + self.CHUNK_SIZE // 2, end),
                )
                if boundary > start:
                    end = boundary + 1
            content = cleaned[start:end].strip()
            if content:
                chunks.append({**metadata, "chunk_index": index, "text": content})
                index += 1
            if end >= len(cleaned):
                break
            start = max(start + 1, end - self.CHUNK_OVERLAP)
        return chunks


class VectorStore:
    def __init__(
        self,
        model_name: str = "sentence-transformers/all-MiniLM-L6-v2",
        embedding_model=None,
        faiss_module=None,
    ):
        self.model_name = model_name
        self.embedding_model = embedding_model
        self._faiss = faiss_module
        self.index = None
        self.documents: list[dict] = []
        self.dimension = 384
        self._lock = threading.RLock()

    def _ensure_backend(self) -> None:
        if self.embedding_model is None:
            try:
                from sentence_transformers import SentenceTransformer
            except ImportError as exc:
                raise RuntimeError("缺少 sentence-transformers") from exc
            self.embedding_model = SentenceTransformer(self.model_name)
        if self._faiss is None:
            try:
                import faiss
            except ImportError as exc:
                raise RuntimeError("缺少 faiss-cpu") from exc
            self._faiss = faiss
        get_dimension = getattr(self.embedding_model, "get_embedding_dimension", None)
        if get_dimension is None:
            get_dimension = self.embedding_model.get_sentence_embedding_dimension
        dimension = int(get_dimension())
        if dimension != 384:
            raise ValueError(f"嵌入模型维度应为 384，实际为 {dimension}")
        self.dimension = dimension
        if self.index is None:
            self.index = self._faiss.IndexFlatL2(self.dimension)

    def add_documents(self, documents: list[dict]) -> int:
        valid = [doc.copy() for doc in documents if doc.get("text", "").strip()]
        if not valid:
            return 0
        with self._lock:
            self._ensure_backend()
            embeddings = self.embedding_model.encode(
                [doc["text"] for doc in valid],
                convert_to_numpy=True,
                normalize_embeddings=True,
                show_progress_bar=False,
            ).astype("float32")
            self.index.add(embeddings)
            self.documents.extend(valid)
            return len(valid)

    def search(self, query: str, k: int = 3) -> list[dict]:
        if not query.strip() or k <= 0:
            return []
        with self._lock:
            if not self.documents:
                return []
            self._ensure_backend()
            vector = self.embedding_model.encode(
                [query],
                convert_to_numpy=True,
                normalize_embeddings=True,
                show_progress_bar=False,
            ).astype("float32")
            distances, indices = self.index.search(vector, min(k, len(self.documents)))
            results = []
            for distance, index in zip(distances[0], indices[0]):
                if index < 0:
                    continue
                results.append({**self.documents[int(index)], "distance": float(distance)})
            return results

    def clear(self) -> None:
        with self._lock:
            self.documents.clear()
            self.index = None

    def save(self, directory: str | Path) -> None:
        with self._lock:
            self._ensure_backend()
            directory = Path(directory)
            directory.mkdir(parents=True, exist_ok=True)
            # FAISS' Windows file API cannot open paths containing CJK characters.
            # Serializing in memory lets pathlib handle the actual filesystem I/O.
            serialized = self._faiss.serialize_index(self.index)
            (directory / "index.faiss").write_bytes(serialized.tobytes())
            (directory / "documents.json").write_text(
                json.dumps(self.documents, ensure_ascii=False, indent=2), encoding="utf-8"
            )

    def load(self, directory: str | Path) -> int:
        with self._lock:
            directory = Path(directory)
            index_path = directory / "index.faiss"
            docs_path = directory / "documents.json"
            if not index_path.exists() or not docs_path.exists():
                return 0
            self._ensure_backend()
            import numpy as np

            serialized = np.frombuffer(index_path.read_bytes(), dtype="uint8")
            index = self._faiss.deserialize_index(serialized)
            documents = json.loads(docs_path.read_text(encoding="utf-8"))
            if not isinstance(documents, list) or any(
                not isinstance(doc, dict) or not isinstance(doc.get("text"), str)
                for doc in documents
            ):
                raise ValueError("知识库文档元数据格式无效")
            if index.ntotal != len(documents):
                raise ValueError("FAISS 索引与文档元数据数量不一致")
            if index.d != self.dimension:
                raise ValueError("FAISS 索引与嵌入模型维度不一致")
            self.index = index
            self.documents = documents
            return len(self.documents)
