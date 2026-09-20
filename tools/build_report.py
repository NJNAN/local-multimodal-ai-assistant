from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "outputs" / "智能多模态AI助手_实验报告.docx"
ACCENT = "16697A"
NAVY = "16324F"
LIGHT = "EAF3F5"
GRAY = "F3F5F7"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, *, bold=False, color="263238", size=9.5) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def add_title(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True


def add_body(doc: Document, text: str, *, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(6)
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        r.bold = True
        p.add_run(text[len(bold_lead) :])
    else:
        p.add_run(text)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths=None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, color="FFFFFF")
        set_cell_shading(table.rows[0].cells[index], NAVY)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(values):
            set_cell_text(cells[col_index], str(value))
            if row_index % 2:
                set_cell_shading(cells[col_index], GRAY)
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_architecture(doc: Document) -> None:
    table = doc.add_table(rows=5, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    rows = [
        ("输入层", "麦克风 / 键盘", "摄像头"),
        ("感知层", "VAD → SenseVoice", "两帧缓冲 → MediaPipe"),
        ("路由层", "文本 / RAG", "视觉理解 / 手势事件"),
        ("模型层", "Qwen3.5-4B + MiniLM + FAISS", "Qwen3-VL-4B + Pose/Hands/Face"),
        ("交互层", "PyQt5 流式对话", "Edge TTS / 动作执行"),
    ]
    for r_idx, values in enumerate(rows):
        for c_idx, value in enumerate(values):
            set_cell_text(
                table.cell(r_idx, c_idx),
                value,
                bold=c_idx == 0,
                color="FFFFFF" if c_idx == 0 else "263238",
                size=10,
            )
            set_cell_shading(table.cell(r_idx, c_idx), ACCENT if c_idx == 0 else (LIGHT if r_idx % 2 == 0 else "FFFFFF"))
    doc.add_paragraph("图 1  系统分层与数据流（可编辑表格）").alignment = WD_ALIGN_PARAGRAPH.CENTER


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("263238")
    for style_name, size, color in (
        ("Title", 28, NAVY),
        ("Heading 1", 18, NAVY),
        ("Heading 2", 14, ACCENT),
        ("Heading 3", 11.5, NAVY),
    ):
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.35)
        section.right_margin = Cm(2.35)
        section.header_distance = Cm(1)
        section.footer_distance = Cm(1)


def add_header_footer(doc: Document) -> None:
    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.text = "行业工程实践 · 模块二 · 智能多模态 AI 助手"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in header.runs:
            run.font.name = "Microsoft YaHei"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string("607D8B")
        add_page_number(section.footer.paragraphs[0])


def read_results() -> list[dict]:
    with (ROOT / "outputs" / "e2e_test_results.csv").open(encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def build() -> None:
    doc = Document()
    configure_styles(doc)

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(70)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("行业工程实践")
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor.from_string(ACCENT)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run("智能多模态 AI 助手")
    r.bold = True
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r.font.size = Pt(32)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("实验报告 · 完整项目交付版")
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor.from_string("607D8B")
    doc.add_paragraph("\n")
    add_table(doc, ["项目", "内容"], [
        ["课程模块", "02_multimodal_ai_assistant"],
        ["技术路线", "GPU PyTorch · llama.cpp (GGUF) · RAG · 多模态感知"],
        ["验收日期", "2026 年 9 月 19 日"],
        ["姓名 / 学号", "____________________________"],
    ], [4.2, 10.8])
    p = doc.add_paragraph("本报告依据六阶段任务书、真实运行日志与自动验收结果编制。")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(40)
    p.runs[0].font.color.rgb = RGBColor.from_string("607D8B")
    p.add_run().add_break(WD_BREAK.PAGE)

    add_title(doc, "摘要", 1)
    add_body(doc, "本项目设计并实现了一套可在个人电脑本地运行的智能多模态 AI 助手。系统同时接入麦克风、摄像头与键盘输入，利用 WebRTC VAD 进行流式语音切分，以 SenseVoice 完成中英文语音识别，通过 llama.cpp（llama-server）本地运行 Qwen3.5-4B 与 Qwen3-VL-4B，结合 MiniLM 和 FAISS 构建大学养老知识库 RAG，并使用 MediaPipe 实现姿态、人脸与六类手势感知。最终系统以 PyQt5 提供统一界面，并通过流式文本与 Edge TTS 给出多通道反馈。")
    add_body(doc, "本机验证采用 Python 3.10.11、PyTorch 2.11.0+cu128 和 NVIDIA GeForce RTX 4060 Laptop GPU。环境检查 21/21 通过，单元测试 21/21 通过，自动端到端用例 10/10 通过。30 分钟持续运行与真人逐手势展示保留为现场人工验收，不在报告中虚构结果。")
    p = doc.add_paragraph()
    p.add_run("关键词：").bold = True
    p.add_run("多模态交互；语音识别；视觉语言模型；RAG；手势识别；本地大模型")

    add_title(doc, "目录", 1)
    for item in [
        "1 项目目标与需求分析", "2 总体方案与系统架构", "3 开发环境与模型配置",
        "4 六阶段实现过程", "5 关键模块设计", "6 实验设计与结果",
        "7 问题定位与工程改进", "8 启动、演示与验收", "9 总结与展望",
    ]:
        doc.add_paragraph(item)
    doc.add_page_break()

    add_title(doc, "1 项目目标与需求分析", 1)
    add_body(doc, "项目目标是完成一个能够“听、看、检索、思考、表达和执行动作”的桌面 AI 助手。任务书要求不是单一模型调用，而是六个阶段最终形成可运行、可测试、可汇报的完整系统。")
    add_table(doc, ["能力域", "输入", "处理", "输出"], [
        ["语音交互", "16 kHz 麦克风流", "VAD + SenseVoice", "识别文本"],
        ["文本问答", "键盘/ASR 文本", "Qwen3.5 流式生成", "逐字回答"],
        ["知识问答", "养老领域问题", "MiniLM + FAISS + Qwen3.5", "有依据的答案"],
        ["视觉理解", "最近两帧", "Qwen3-VL-4B（llama.cpp）", "场景描述/视觉回答"],
        ["人体与手势", "摄像头画面", "MediaPipe + 规则识别", "六类动作事件"],
        ["语音反馈", "回答文本", "Edge TTS + 播放队列", "中英文语音"],
    ])

    add_title(doc, "2 总体方案与系统架构", 1)
    add_body(doc, "系统采用“采集—感知—路由—模型—交互”五层结构。音频与视频各自在线程中持续采集，CPU 负责轻量感知和界面，GPU 负责 PyTorch ASR 与文本嵌入；llama.cpp（llama-server）按需加载并管理文本与视觉大模型，空闲自动卸载。模态路由器依据关键词及手动命令选择文本、RAG 或视觉链路，避免每次提问都调用较慢的视觉模型。")
    add_architecture(doc)
    add_body(doc, "数据流保持松耦合：工作线程只产生事件，通过 Bridge 信号安全传递到 GUI 主线程；回答生成与 TTS 播放均可取消。知识库索引持久化后在启动时自动加载，避免重复编码 31 份文档。")

    add_title(doc, "3 开发环境与模型配置", 1)
    add_table(doc, ["组件", "版本/模型", "运行位置"], [
        ["Python", "3.10.11", "D:\\mm_ai_env"],
        ["PyTorch", "2.11.0+cu128", "RTX 4060 / CUDA"],
        ["llama.cpp", "llama-server b10689（Vulkan 构建）", "本地进程 / GPU"],
        ["文本模型", "Qwen3.5-4B IQ4_XS（GGUF，2.55 GB）", "llama.cpp · GPU 全量 offload"],
        ["视觉模型", "Qwen3-VL-4B-Instruct Q4_K_M + mmproj Q8_0", "llama.cpp · GPU（显存 < 8 GB）"],
        ["ASR", "SenseVoiceSmall", "PyTorch CUDA"],
        ["嵌入", "all-MiniLM-L6-v2, 384 维", "SentenceTransformers CUDA"],
        ["感知", "MediaPipe 0.10.14", "CPU / TFLite"],
        ["界面与播报", "PyQt5 + Edge TTS + pygame", "桌面端"],
    ])
    add_body(doc, "依赖安装使用国内镜像：常规 Python 包采用清华 PyPI，CUDA 版 PyTorch 采用上海交大 PyTorch 镜像；SenseVoice 从 ModelScope 国内源下载。文本与视觉 GGUF 经 HuggingFace/ModelScope 下载并完成 SHA256 校验。最终所有模型均已落盘，正常启动不再产生下载量。")

    add_title(doc, "4 六阶段实现过程", 1)
    phases = [
        ("4.1 阶段一：环境与音视频采集", "统一配置目录、日志和随机种子；完成 PyAudio 音频采集、OpenCV 视频采集、帧缓冲与采集测试脚本。摄像头实测输出 640×480 JPEG。"),
        ("4.2 阶段二：VAD 与流式 ASR", "按 30 ms 帧处理语音：0.6 s 静音断句、0.5 s 最短语音、240 ms 前滚和 120 ms 尾部。AudioStreamThread 将采集、VAD、识别和回调解耦。SenseVoice 在 cuda:0 上完成合成测试音频识别。"),
        ("4.3 阶段三：文本对话与知识库 RAG", "通过 llama-server 的 OpenAI 兼容 HTTP API 对接 Qwen3.5-4B，支持流式生成、取消和 8192 上下文。知识库新增 DOCX 解析以适配随项目提供的 31 份养老资料，共生成 53 个知识块并建立 FAISS 索引。"),
        ("4.4 阶段四：视觉理解与语音合成", "视觉路由从视频缓冲选择最近两帧，经 JPEG/Base64 送入 Qwen3-VL-4B（llama.cpp 多模态，mmproj Q8_0）。TTS 自动区分中英文声音，StreamingTTSPlayer 使用生产者—消费者队列避免阻塞界面。"),
        ("4.5 阶段五：姿态与手势识别", "HumanDetector 并行处理 Pose、Hands 和 Face；GestureRecognizer 识别点赞、点踩、张掌、食指向上、食指向下和 V 字，并通过连续帧确认、冷却时间减少抖动。"),
        ("4.6 阶段六：系统联调与答辩", "AssistantService 统一调度文本、RAG、视觉与取消逻辑；PyQt5 界面展示视频、聊天、知识库导入、视觉分析和停止控制。验收脚本导出 CSV 指标，报告和 PPT 与代码一并交付。"),
    ]
    for title, text in phases:
        add_title(doc, title, 2)
        add_body(doc, text)

    add_title(doc, "5 关键模块设计", 1)
    add_title(doc, "5.1 VAD 参数与字节计算", 2)
    add_body(doc, "在 16 kHz、16 bit、单声道条件下，30 ms 包含 16,000 × 0.03 = 480 个采样点；每个采样点 2 字节，因此帧长度为 960 字节。实现中对帧长严格校验，避免任务书示例中把“采样点数”误写成“字节数”导致 WebRTC VAD 报错。")
    add_title(doc, "5.2 RAG 持久化与中文路径", 2)
    add_body(doc, "文档解析后按 500 字符分块并保留 50 字符重叠。嵌入向量先归一化，再进入 IndexFlatL2。由于 FAISS 的 Windows 文件 API 无法直接写入含中文路径，系统改用 serialize_index / deserialize_index 在内存中转换，由 pathlib 完成实际读写。")
    add_title(doc, "5.3 模态路由与线程安全", 2)
    add_body(doc, "包含“看看、画面、摄像头”等视觉意图的请求进入视觉模式；知识库导入后可进入 RAG；其余走文本模式。工作线程从不直接修改控件，而是通过 Qt 信号回到主线程。停止按钮同时取消 LLM 生成并清空 TTS 队列。")
    add_title(doc, "5.4 手势—动作映射", 2)
    add_table(doc, ["手势", "内部标识", "动作"], [
        ["点赞", "thumbs_up", "增大音量"], ["点踩", "thumbs_down", "减小音量"],
        ["张掌", "open_palm", "停止生成与播报"], ["食指向上", "index_up", "开始聆听"],
        ["食指向下", "index_down", "静音"], ["V 字", "victory", "保存截图"],
    ])

    add_title(doc, "6 实验设计与结果", 1)
    add_body(doc, "测试分为单元测试、环境检查、模块实测、端到端测试和 GUI 冒烟测试。所有性能数字均来自 2026-09-13 的本机运行，不使用估算值替代实测。")
    results = read_results()
    add_table(doc, ["编号", "场景", "状态", "延迟/s"], [
        [r["case_id"], r["scenario"], r["status"], r["latency_seconds"]]
        for r in results
    ], [2.1, 7.4, 2.3, 2.5])
    add_bullets(doc, [
        "环境检查：21/21 通过；CUDA 可用，GPU 为 RTX 4060 Laptop GPU。",
        "单元测试：21/21 通过，覆盖 VAD、ASR 清洗、路由/TTS、LLM（llama.cpp 后端）、知识解析和手势逻辑。",
        "llama.cpp GPU 实测：文本首 token 0.157 s、生成 39.5~67.8 tokens/s；视觉单图理解 2.27 s；模型加载约 6.2 s。",
        "SenseVoice 端到端用例包含模型加载；单段前向 RTF 实测约 0.13。",
        "GUI 冒烟测试：创建全部组件、模型→TTS 链路合成 MP3、安全退出无残留进程。",
        "进程树 RSS：文本约 3.2 GB / 视觉约 2.6 GB（llama-server 含在内）。",
    ])

    add_title(doc, "7 问题定位与工程改进", 1)
    add_table(doc, ["问题", "根因", "解决方案"], [
        ["MediaPipe 初始化找不到资源", "Windows 原生库不能处理中文虚拟环境路径", "固定使用 D:\\mm_ai_env，并补齐 heavy pose 模型"],
        ["SenseVoice 词表加载失败", "SentencePiece 原生库不能处理中文模型路径", "模型放到 D:\\mm_ai_models\\SenseVoiceSmall"],
        ["FAISS 索引无法保存", "原生文件 API 不支持中文路径", "改用内存序列化 + pathlib"],
        ["llama-cpp-python 缺 Qwen3-VL 支持", "0.3.35 无新模型专用处理器", "改用官方 llama-server（OpenAI 兼容 API），上层接口不变"],
        ["8 GB 显存装不下双模型常驻", "文本+视觉同时驻留会超上限", "独占式生命周期 + --fit 自动 offload，空闲自动卸载"],
        ["GGUF 跨境下载慢/断流", "5.5 GB 大文件与跨境链路波动", "ModelScope 镜像 + aria2 断点续传 + SHA256 校验"],
        ["语音连发导致请求堆积与模型来回切换", "每段语音独立提问且无排队；“看”字误匹配触发视觉", "单槽队列保留最新一条；视觉关键词精化；模型加载状态提示"],
        ["Qt 后台线程更新按钮", "跨线程控件访问不安全", "新增 knowledge_imported 信号回主线程"],
    ])

    add_title(doc, "8 启动、演示与验收", 1)
    add_title(doc, "8.1 启动方式", 2)
    add_body(doc, "推荐双击项目根目录的“启动助手.bat”。模型由 llama.cpp 按需自动加载，无需预先启动任何服务；脚本先执行 21 项环境检查再打开 GUI。手动命令如下：")
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    r = p.add_run('cd "D:\\作业\\行业工程实践\\projects\\02_multimodal_ai_assistant"\nD:\\mm_ai_env\\Scripts\\python.exe scripts\\00_check_env.py\nD:\\mm_ai_env\\Scripts\\python.exe main.py')
    r.font.name = "Consolas"
    r.font.size = Pt(9)
    set_cell_shading if False else None
    add_title(doc, "8.2 建议的 10 分钟演示顺序", 2)
    add_bullets(doc, [
        "第 1 分钟：运行环境检查，说明 GPU PyTorch 与两个 llama.cpp 本地模型（Qwen3.5-4B / Qwen3-VL-4B）。",
        "第 2–3 分钟：语音提问，展示 0.6 秒停顿自动断句与流式回答。",
        "第 4–5 分钟：询问大学养老问题，对比知识库回答与普通回答。",
        "第 6 分钟：点击“分析画面”，展示两帧视觉理解。",
        "第 7–8 分钟：依次做六种手势，展示动作映射与冷却机制。",
        "第 9 分钟：展示中英文 TTS、停止生成和安全退出。",
        "第 10 分钟：打开 CSV 指标与实验报告，总结工程问题及解决方案。",
    ])
    add_title(doc, "8.3 人工验收提醒", 2)
    add_body(doc, "30 分钟连续运行应在答辩或验收前实际执行并保留日志；六种手势应由本人在摄像头前逐项展示。自动测试验证了识别规则、映射和防抖逻辑，但不把真人准确率作为未实测数据写入报告。")

    add_title(doc, "9 总结与展望", 1)
    add_body(doc, "本项目已经从单模块样例发展为可启动、可交互、可复测的完整桌面多模态系统。其工程价值在于把不同运行时和不同延迟特性的模型组织为稳定的数据流，并针对 Windows 中文路径、GPU 依赖与本地服务代理问题给出可复用的解决办法。2026 年 9 月进一步将推理后端从 Ollama 迁移到 llama.cpp（Qwen3.5-4B / Qwen3-VL-4B GGUF，SHA256 校验、按需加载、--fit 自动 offload），显著提升了吞吐与响应速度。")
    add_body(doc, "后续可进一步量化真实环境下的手势混淆矩阵、端到端首字延迟和 30 分钟资源曲线，并增加完全离线的本地 TTS，以降低对网络语音服务的依赖。")
    add_title(doc, "附录 A：主要交付文件", 1)
    add_bullets(doc, [
        "main.py：系统入口；启动助手.bat / 启动助手.ps1：一键启动。",
        "src/：采集、VAD、ASR、LLM、RAG、视觉、TTS、感知、GUI 与服务层。",
        "scripts/00–12：从环境检查到端到端验收的独立测试入口（含性能采集与链路检查）。",
        "outputs/e2e_test_results.csv 与 performance_metrics.csv：真实验收数据。",
        "docs/架构与验收记录.md、docs/模型替换验收报告.md：架构、关键设计与 llama.cpp 替换验收记录。",
    ])

    configure_styles(doc)
    add_header_footer(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
